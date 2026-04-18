"""Router del Portal de Cliente."""
import io
import os
import json
from datetime import datetime
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
import models, schemas
from .deps import verify_token

def _borrar_docs_asociados(db: Session, contextos: list[str], contexto_id: int):
    """Borra documentos de R2/disco y DB asociados a una solicitud."""
    from .documentos import _delete_file
    docs = db.query(models.Documento).filter(
        models.Documento.contexto.in_(contextos),
        models.Documento.contexto_id == contexto_id,
    ).all()
    for d in docs:
        _delete_file(d.ruta)
        db.delete(d)
    return len(docs)

router = APIRouter(prefix="/portal", tags=["portal"])


def _require_portal(token=Depends(verify_token)):
    if token.get("rol") not in ("admin", "cliente"):
        raise HTTPException(status_code=403, detail="Acceso solo para clientes o administradores")
    return token


# ─── AFILIADOS ────────────────────────────────────────────────────────────────

@router.get("/afiliados")
def portal_afiliados(q: str = "", db: Session = Depends(get_db), token=Depends(_require_portal)):
    rol = token.get("rol", "")
    cliente_ref = (token.get("cliente_ref") or "").strip()

    # Cliente sin cliente_ref configurado → no mostrar nada
    if rol != "admin" and not cliente_ref:
        raise HTTPException(400, "Este usuario no tiene un cliente asociado. Contacte al administrador.")

    query = db.query(models.Afiliado).filter(models.Afiliado.activo == True)
    if rol != "admin":
        # Filtro estricto: solo afiliados cuyo cliente_txt coincide exactamente
        query = query.filter(
            models.Afiliado.cliente_txt == cliente_ref,
            models.Afiliado.cliente_txt != None,
            models.Afiliado.cliente_txt != "",
        )
    if q:
        from sqlalchemy import or_
        query = query.filter(
            or_(models.Afiliado.nombre.ilike(f"%{q}%"), models.Afiliado.doc.ilike(f"%{q}%"))
        )
    afiliados = query.order_by(models.Afiliado.nombre).limit(1000).all()

    result = []
    for a in afiliados:
        try:
            srvs = json.loads(a.servicios or "[]")
        except Exception:
            srvs = []
        result.append({
            "id": a.id, "nombre": a.nombre, "doc": a.doc, "tipo_doc": a.tipo_doc,
            "empresa": a.empresa, "cargo": a.cargo,
            "eps": a.eps, "afp": a.afp, "ccf": a.ccf, "arl": a.arl,
            "estado": a.estado, "estado_srv": a.estado_srv,
            "servicios": srvs, "tel": a.tel, "email": a.email,
            "fecha_ingreso": a.fecha_ingreso, "fecha_afiliacion": a.fecha_afiliacion,
            "novedades": a.novedades or "",
            "detalle": a.detalle or "",
        })
    return result


@router.get("/afiliados/{doc}/resumen")
def portal_resumen(doc: str, db: Session = Depends(get_db), token=Depends(_require_portal)):
    rol = token.get("rol", "")
    cliente_ref = (token.get("cliente_ref") or "").strip()
    afil = db.query(models.Afiliado).filter_by(doc=doc, activo=True).first()
    if not afil:
        raise HTTPException(404, "Afiliado no encontrado")
    if rol != "admin" and afil.cliente_txt != cliente_ref:
        raise HTTPException(403, "Sin acceso a este afiliado")

    facturas = (db.query(models.Factura)
                .filter_by(doc=doc, afiliado_eliminado=False)
                .order_by(models.Factura.anio.desc(), models.Factura.id.desc())
                .all())

    try:
        srvs = json.loads(afil.servicios or "[]")
    except Exception:
        srvs = []

    return {
        "afiliado": {
            "id": afil.id, "nombre": afil.nombre, "doc": afil.doc, "tipo_doc": afil.tipo_doc,
            "empresa": afil.empresa, "cargo": afil.cargo,
            "eps": afil.eps, "afp": afil.afp, "ccf": afil.ccf, "arl": afil.arl,
            "estado": afil.estado, "estado_srv": afil.estado_srv,
            "servicios": srvs, "tel": afil.tel, "email": afil.email,
            "dir": afil.dir, "obs": afil.obs, "ibc": afil.ibc,
            "fecha_ingreso": afil.fecha_ingreso, "fecha_afiliacion": afil.fecha_afiliacion,
        },
        "facturas": [{
            "id": f.id, "codigo": f.codigo, "mes": f.mes, "anio": f.anio,
            "estado": f.estado, "costos": f.costos, "banco": f.banco,
            "pagado_en": f.pagado_en.isoformat() if f.pagado_en else None,
        } for f in facturas],
        "total_pendiente": sum(f.costos or 0 for f in facturas if f.estado == "pendiente"),
        "total_pagado":    sum(f.costos or 0 for f in facturas if f.estado in ("pagado", "planilla_pagada")),
    }


# ─── NOVEDADES DE PAGO ────────────────────────────────────────────────────────

@router.post("/novedades-pago", status_code=201)
def portal_novedad_pago(
    data: schemas.NovedadPagoCreate,
    db: Session = Depends(get_db),
    token=Depends(_require_portal),
):
    rol             = token.get("rol", "")
    cliente_ref     = (token.get("cliente_ref") or "").strip()
    username        = token.get("sub", "")
    nombre_cliente  = token.get("nombre", username)

    if not data.afiliados_docs:
        raise HTTPException(400, "Debe seleccionar al menos un afiliado")

    afiliados = []
    for doc in data.afiliados_docs:
        afil = db.query(models.Afiliado).filter_by(doc=doc, activo=True).first()
        if not afil:
            raise HTTPException(404, f"Afiliado con documento {doc} no encontrado")
        if rol != "admin" and afil.cliente_txt != cliente_ref:
            raise HTTPException(403, f"Sin acceso al afiliado con documento {doc}")
        afiliados.append(afil)

    novedad = models.NovedadPago(
        cliente_ref=cliente_ref,
        username_cliente=username,
        mes=data.mes,
        anio=data.anio,
        afiliados_docs=json.dumps(data.afiliados_docs),
        afiliados_nombres=json.dumps([a.nombre for a in afiliados]),
        obs=data.obs,
        estado="pendiente",
    )
    db.add(novedad)

    admins = db.query(models.Usuario).filter_by(rol="admin", activo=True).all()
    for admin in admins:
        db.add(models.Notificacion(
            usuario=admin.username,
            mensaje=f"Cliente '{nombre_cliente}' reportó novedad SS para {len(afiliados)} afiliado(s) — {data.mes} {data.anio}",
        ))

    db.commit()
    db.refresh(novedad)
    return {"ok": True, "id": novedad.id}


@router.get("/novedades-pago")
def portal_list_novedades(db: Session = Depends(get_db), token=Depends(_require_portal)):
    query = db.query(models.NovedadPago)
    if token.get("rol") != "admin":
        query = query.filter_by(username_cliente=token.get("sub"))
    rows = query.order_by(models.NovedadPago.id.desc()).limit(500).all()
    return [{
        "id": r.id, "cliente_ref": r.cliente_ref, "username_cliente": r.username_cliente,
        "mes": r.mes, "anio": r.anio,
        "afiliados": json.loads(r.afiliados_nombres or "[]"),
        "obs": r.obs, "estado": r.estado, "respuesta": r.respuesta or "",
        "creado": r.creado.isoformat() if r.creado else None,
    } for r in rows]


@router.get("/novedades-pago/{id}/exportar-excel")
def exportar_novedad_pago_excel(id: int, db: Session = Depends(get_db), token=Depends(verify_token)):
    """Admin descarga Excel con la info completa de los afiliados reportados en una novedad de pago."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    if token.get("rol") != "admin":
        raise HTTPException(403, "Solo administradores pueden exportar")

    nov = db.query(models.NovedadPago).filter_by(id=id).first()
    if not nov:
        raise HTTPException(404, "Novedad no encontrada")

    docs = json.loads(nov.afiliados_docs or "[]")
    afiliados = db.query(models.Afiliado).filter(models.Afiliado.doc.in_(docs)).all()
    # Ordenar según el orden original reportado por el cliente
    orden = {doc: i for i, doc in enumerate(docs)}
    afiliados.sort(key=lambda a: orden.get(a.doc, 999))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Novedad {nov.mes} {nov.anio}"

    header_fill = PatternFill("solid", fgColor="0D3B6E")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    headers = ["Nombre", "Tipo Doc", "Documento", "Empresa", "Cargo",
               "Estado", "EPS", "AFP", "CCF", "ARL",
               "Teléfono", "Email", "IBC"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, a in enumerate(afiliados, 2):
        ws.cell(row=row, column=1,  value=a.nombre)
        ws.cell(row=row, column=2,  value=a.tipo_doc)
        ws.cell(row=row, column=3,  value=a.doc)
        ws.cell(row=row, column=4,  value=a.empresa)
        ws.cell(row=row, column=5,  value=a.cargo)
        ws.cell(row=row, column=6,  value=a.estado)
        ws.cell(row=row, column=7,  value=a.eps)
        ws.cell(row=row, column=8,  value=a.afp)
        ws.cell(row=row, column=9,  value=a.ccf)
        ws.cell(row=row, column=10, value=a.arl)
        ws.cell(row=row, column=11, value=a.tel)
        ws.cell(row=row, column=12, value=a.email)
        ws.cell(row=row, column=13, value=a.ibc)

    # Fila de resumen al final
    if afiliados:
        ws.append([])
        ws.append([f"Novedad de Pago — {nov.mes} {nov.anio}",
                   f"Cliente: {nov.cliente_ref or nov.username_cliente}",
                   f"Total afiliados: {len(afiliados)}",
                   f"Obs: {nov.obs or '—'}"])

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    nombre_archivo = f"novedad-pago-{nov.mes}-{nov.anio}-{id}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )


@router.patch("/novedades-pago/{id}/estado")
def portal_update_novedad_estado(
    id: int, body: schemas.EstadoSolicitudBody,
    db: Session = Depends(get_db),
    token=Depends(verify_token),
):
    if token.get("rol") != "admin":
        raise HTTPException(403, "Solo administradores pueden actualizar el estado")
    nov = db.query(models.NovedadPago).filter_by(id=id).first()
    if not nov:
        raise HTTPException(404, "Novedad no encontrada")
    nov.estado = body.estado
    if body.respuesta is not None:
        nov.respuesta = body.respuesta
    if nov.username_cliente:
        msg = f"Tu novedad de pago ({nov.mes} {nov.anio}) fue marcada como '{body.estado}'"
        if nov.respuesta:
            msg += f". Nota del administrador: {nov.respuesta}"
        db.add(models.Notificacion(usuario=nov.username_cliente, mensaje=msg))
    db.commit()
    return {"ok": True}


@router.delete("/novedades-pago/{id}")
def portal_delete_novedad(id: int, db: Session = Depends(get_db), token=Depends(verify_token)):
    if token.get("rol") != "admin":
        raise HTTPException(403, "Solo administradores pueden eliminar novedades")
    nov = db.query(models.NovedadPago).filter_by(id=id).first()
    if not nov:
        raise HTTPException(404, "Novedad no encontrada")
    _borrar_docs_asociados(db, ['novedad_pago', 'resp_pago'], id)
    db.delete(nov)
    db.commit()
    return {"ok": True}


# ─── SOLICITUDES DE RETIRO ────────────────────────────────────────────────────

@router.post("/solicitar-retiro", status_code=201)
def portal_solicitar_retiro(
    data: schemas.SolicitudRetiroCreate,
    db: Session = Depends(get_db),
    token=Depends(_require_portal),
):
    rol            = token.get("rol", "")
    cliente_ref    = (token.get("cliente_ref") or "").strip()
    username       = token.get("sub", "")
    nombre_cliente = token.get("nombre", username)

    afil = db.query(models.Afiliado).filter_by(doc=data.afiliado_doc, activo=True).first()
    if not afil:
        raise HTTPException(404, "Afiliado no encontrado")
    if rol != "admin" and afil.cliente_txt != cliente_ref:
        raise HTTPException(403, "Sin acceso a este afiliado")

    solicitud = models.SolicitudRetiro(
        cliente_ref=cliente_ref,
        username_cliente=username,
        afiliado_doc=data.afiliado_doc,
        afiliado_nombre=afil.nombre,
        motivo=data.motivo,
        obs=data.obs,
        estado="pendiente",
    )
    db.add(solicitud)

    admins = db.query(models.Usuario).filter_by(rol="admin", activo=True).all()
    for admin in admins:
        db.add(models.Notificacion(
            usuario=admin.username,
            mensaje=f"Cliente '{nombre_cliente}' solicita retiro de '{afil.nombre}' — {data.motivo}",
        ))

    db.commit()
    db.refresh(solicitud)
    return {"ok": True, "id": solicitud.id}


@router.get("/solicitudes-retiro")
def portal_list_solicitudes(db: Session = Depends(get_db), token=Depends(_require_portal)):
    query = db.query(models.SolicitudRetiro)
    if token.get("rol") != "admin":
        query = query.filter_by(username_cliente=token.get("sub"))
    rows = query.order_by(models.SolicitudRetiro.id.desc()).limit(500).all()
    return [{
        "id": r.id, "cliente_ref": r.cliente_ref,
        "afiliado_doc": r.afiliado_doc, "afiliado_nombre": r.afiliado_nombre,
        "motivo": r.motivo, "obs": r.obs, "estado": r.estado, "respuesta": r.respuesta or "",
        "creado": r.creado.isoformat() if r.creado else None,
    } for r in rows]


@router.patch("/solicitudes-retiro/{id}/estado")
def portal_update_solicitud_estado(
    id: int, body: schemas.EstadoSolicitudBody,
    db: Session = Depends(get_db),
    token=Depends(verify_token),
):
    if token.get("rol") != "admin":
        raise HTTPException(403, "Solo administradores pueden actualizar el estado")
    sol = db.query(models.SolicitudRetiro).filter_by(id=id).first()
    if not sol:
        raise HTTPException(404, "Solicitud no encontrada")
    sol.estado = body.estado
    if body.respuesta is not None:
        sol.respuesta = body.respuesta
    if sol.username_cliente:
        msg = f"Tu solicitud de retiro de '{sol.afiliado_nombre}' fue marcada como '{body.estado}'"
        if sol.respuesta:
            msg += f". Nota del administrador: {sol.respuesta}"
        db.add(models.Notificacion(usuario=sol.username_cliente, mensaje=msg))
    db.commit()
    return {"ok": True}


@router.delete("/solicitudes-retiro/{id}")
def portal_delete_solicitud_retiro(id: int, db: Session = Depends(get_db), token=Depends(verify_token)):
    if token.get("rol") != "admin":
        raise HTTPException(403, "Solo administradores pueden eliminar solicitudes")
    sol = db.query(models.SolicitudRetiro).filter_by(id=id).first()
    if not sol:
        raise HTTPException(404, "Solicitud no encontrada")
    _borrar_docs_asociados(db, ['novedad_retiro', 'resp_retiro'], id)
    db.delete(sol)
    db.commit()
    return {"ok": True}


# ─── SOLICITUDES DE NOVEDAD ───────────────────────────────────────────────────

@router.post("/solicitudes-novedad", status_code=201)
def portal_crear_novedad(
    data: schemas.SolicitudNovedadCreate,
    db: Session = Depends(get_db),
    token=Depends(_require_portal),
):
    rol            = token.get("rol", "")
    cliente_ref    = (token.get("cliente_ref") or "").strip()
    username       = token.get("sub", "")
    nombre_cliente = token.get("nombre", username)

    afil = db.query(models.Afiliado).filter_by(doc=data.afiliado_doc, activo=True).first()
    if not afil:
        raise HTTPException(404, "Afiliado no encontrado")
    if rol != "admin" and afil.cliente_txt != cliente_ref:
        raise HTTPException(403, "Sin acceso a este afiliado")

    solicitud = models.SolicitudNovedad(
        cliente_ref=cliente_ref,
        username_cliente=username,
        afiliado_doc=data.afiliado_doc,
        afiliado_nombre=afil.nombre,
        tipo=data.tipo,
        descripcion=data.descripcion,
        estado="pendiente",
    )
    db.add(solicitud)

    admins = db.query(models.Usuario).filter_by(rol="admin", activo=True).all()
    for admin in admins:
        db.add(models.Notificacion(
            usuario=admin.username,
            mensaje=f"Cliente '{nombre_cliente}' reportó novedad '{data.tipo}' para '{afil.nombre}'",
        ))

    db.commit()
    db.refresh(solicitud)
    return {"ok": True, "id": solicitud.id}


@router.get("/solicitudes-novedad")
def portal_list_novedades_afil(db: Session = Depends(get_db), token=Depends(_require_portal)):
    query = db.query(models.SolicitudNovedad)
    if token.get("rol") != "admin":
        query = query.filter_by(username_cliente=token.get("sub"))
    rows = query.order_by(models.SolicitudNovedad.id.desc()).limit(500).all()
    return [{
        "id": r.id, "cliente_ref": r.cliente_ref,
        "afiliado_doc": r.afiliado_doc, "afiliado_nombre": r.afiliado_nombre,
        "tipo": r.tipo, "descripcion": r.descripcion,
        "estado": r.estado, "respuesta": r.respuesta or "",
        "creado": r.creado.isoformat() if r.creado else None,
    } for r in rows]


# ─── EXPORTAR EXCEL ───────────────────────────────────────────────────────────

@router.get("/exportar-excel")
def portal_exportar_excel(db: Session = Depends(get_db), token=Depends(_require_portal)):
    """Exporta los afiliados del cliente como archivo Excel."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    rol = token.get("rol", "")
    cliente_ref = (token.get("cliente_ref") or "").strip()

    if rol != "admin" and not cliente_ref:
        raise HTTPException(400, "Este usuario no tiene un cliente asociado.")

    query = db.query(models.Afiliado).filter(models.Afiliado.activo == True)
    if rol != "admin":
        query = query.filter(
            models.Afiliado.cliente_txt == cliente_ref,
            models.Afiliado.cliente_txt != None,
            models.Afiliado.cliente_txt != "",
        )
    afiliados = query.order_by(models.Afiliado.nombre).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mis Afiliados"

    headers = ["Nombre", "Tipo Doc", "Documento", "Empresa", "Cargo",
               "Estado", "Estado Servicio", "EPS", "AFP", "CCF", "ARL",
               "Teléfono", "Email", "Fecha Ingreso", "Fecha Afiliación"]

    header_fill = PatternFill("solid", fgColor="0D3B6E")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, a in enumerate(afiliados, 2):
        ws.cell(row=row, column=1,  value=a.nombre)
        ws.cell(row=row, column=2,  value=a.tipo_doc)
        ws.cell(row=row, column=3,  value=a.doc)
        ws.cell(row=row, column=4,  value=a.empresa)
        ws.cell(row=row, column=5,  value=a.cargo)
        ws.cell(row=row, column=6,  value=a.estado)
        ws.cell(row=row, column=7,  value=a.estado_srv)
        ws.cell(row=row, column=8,  value=a.eps)
        ws.cell(row=row, column=9,  value=a.afp)
        ws.cell(row=row, column=10, value=a.ccf)
        ws.cell(row=row, column=11, value=a.arl)
        ws.cell(row=row, column=12, value=a.tel)
        ws.cell(row=row, column=13, value=a.email)
        ws.cell(row=row, column=14, value=a.fecha_ingreso)
        ws.cell(row=row, column=15, value=a.fecha_afiliacion)

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="mis-afiliados.xlsx"'},
    )


@router.patch("/solicitudes-novedad/{id}/estado")
def portal_update_novedad_afil_estado(
    id: int, body: schemas.EstadoSolicitudBody,
    db: Session = Depends(get_db),
    token=Depends(verify_token),
):
    if token.get("rol") != "admin":
        raise HTTPException(403, "Solo administradores pueden actualizar el estado")
    s = db.query(models.SolicitudNovedad).filter_by(id=id).first()
    if not s:
        raise HTTPException(404, "Solicitud no encontrada")
    s.estado = body.estado
    if body.respuesta is not None:
        s.respuesta = body.respuesta
    if s.username_cliente:
        msg = f"Tu novedad '{s.tipo}' para '{s.afiliado_nombre}' fue marcada como '{body.estado}'"
        if s.respuesta:
            msg += f". Nota del administrador: {s.respuesta}"
        db.add(models.Notificacion(usuario=s.username_cliente, mensaje=msg))
    db.commit()
    return {"ok": True}


@router.delete("/solicitudes-novedad/{id}")
def portal_delete_novedad_afil(id: int, db: Session = Depends(get_db), token=Depends(verify_token)):
    if token.get("rol") != "admin":
        raise HTTPException(403, "Solo administradores pueden eliminar solicitudes")
    s = db.query(models.SolicitudNovedad).filter_by(id=id).first()
    if not s:
        raise HTTPException(404, "Solicitud no encontrada")
    _borrar_docs_asociados(db, ['novedad_afil', 'resp_afil'], id)
    db.delete(s)
    db.commit()
    return {"ok": True}


# ─── ESTADO DE CUENTA PDF (accesible para clientes) ─────────────────────────

@router.get("/afiliados/{doc}/estado-cuenta")
def portal_estado_cuenta(doc: str, db: Session = Depends(get_db), token=Depends(_require_portal)):
    """Genera PDF de estado de cuenta accesible para clientes."""
    from routers.afiliados import estado_cuenta_afiliado as _gen_pdf

    rol = token.get("rol", "")
    cliente_ref = (token.get("cliente_ref") or "").strip()

    afil = db.query(models.Afiliado).filter_by(doc=doc, activo=True).first()
    if not afil:
        raise HTTPException(404, "Afiliado no encontrado")
    if rol != "admin" and afil.cliente_txt != cliente_ref:
        raise HTTPException(403, "Sin acceso a este afiliado")

    return _gen_pdf(afil.id, db, token)


# ─── PLANILLAS PAGADAS (para portal cliente) ─────────────────────────────────

@router.get("/planillas")
def portal_planillas(mes: str = "", anio: str = "",
                     db: Session = Depends(get_db), token=Depends(_require_portal)):
    """Lista las planillas de pago SS del cliente autenticado."""
    rol = token.get("rol", "")
    cliente_ref = (token.get("cliente_ref") or "").strip()

    q = db.query(models.PlanillaPago).order_by(models.PlanillaPago.id.desc())
    if rol != "admin":
        if not cliente_ref:
            raise HTTPException(400, "Usuario sin cliente asociado")
        q = q.filter(models.PlanillaPago.cliente_ref == cliente_ref)
    if mes:  q = q.filter(models.PlanillaPago.mes == mes)
    if anio: q = q.filter(models.PlanillaPago.anio == anio)

    rows = q.all()
    if not rows:
        return []
    # Batch load all documents for these planillas (avoid N+1)
    planilla_ids = [p.id for p in rows]
    all_docs = db.query(models.Documento).filter(
        models.Documento.contexto == "planilla_pago",
        models.Documento.contexto_id.in_(planilla_ids),
    ).all()
    docs_by_planilla = {}
    for d in all_docs:
        docs_by_planilla.setdefault(d.contexto_id, []).append(d)
    result = []
    for p in rows:
        docs = docs_by_planilla.get(p.id, [])
        result.append({
            "id": p.id, "cliente_ref": p.cliente_ref, "mes": p.mes, "anio": p.anio,
            "observaciones": p.observaciones,
            "creado": p.creado.isoformat() if p.creado else None,
            "archivos": [{"id": d.id, "nombre": d.nombre, "tamano": d.tamano} for d in docs],
        })
    return result


# ─── AVISOS (admin → cliente) ─────────────────────────────────────────────────

@router.post("/avisos", status_code=201)
def crear_aviso(body: schemas.AvisoClienteCreate, db: Session = Depends(get_db), token=Depends(verify_token)):
    """Admin crea un aviso dirigido a un cliente específico."""
    if token.get("rol") != "admin":
        raise HTTPException(status_code=403, detail="Solo administradores pueden crear avisos")
    aviso = models.AvisoCliente(
        cliente_ref=body.cliente_ref.strip(),
        titulo=body.titulo.strip(),
        mensaje=body.mensaje.strip(),
        creado_por=token.get("sub", ""),
    )
    db.add(aviso)
    db.flush()  # get aviso.id before commit
    # Notificar a todos los usuarios del cliente
    usuarios_cliente = db.query(models.Usuario).filter_by(
        cliente_ref=body.cliente_ref.strip(), activo=True
    ).all()
    for u in usuarios_cliente:
        db.add(models.Notificacion(
            usuario=u.username,
            mensaje=f"Nuevo aviso: {body.titulo[:80]}",
            tarea_id=None,
        ))
    db.commit()
    return {"id": aviso.id, "ok": True}


@router.get("/avisos")
def listar_avisos(db: Session = Depends(get_db), token=Depends(_require_portal)):
    """Admin ve todos los avisos; cliente ve solo los suyos."""
    rol = token.get("rol", "")
    cliente_ref = (token.get("cliente_ref") or "").strip()

    q = db.query(models.AvisoCliente).order_by(models.AvisoCliente.id.desc())
    if rol != "admin":
        if not cliente_ref:
            raise HTTPException(400, "Usuario sin cliente asociado")
        q = q.filter(models.AvisoCliente.cliente_ref == cliente_ref)

    rows = q.all()
    if not rows:
        return []
    # Batch load documents (avoid N+1)
    aviso_ids = [a.id for a in rows]
    all_docs = db.query(models.Documento).filter(
        models.Documento.contexto == "aviso",
        models.Documento.contexto_id.in_(aviso_ids),
    ).all()
    docs_by_aviso = {}
    for d in all_docs:
        docs_by_aviso.setdefault(d.contexto_id, []).append(
            {"id": d.id, "nombre": d.nombre, "tipo": d.tipo, "tamano": d.tamano}
        )
    return [
        {
            "id": a.id,
            "cliente_ref": a.cliente_ref,
            "titulo": a.titulo,
            "mensaje": a.mensaje,
            "leido": a.leido,
            "creado_por": a.creado_por,
            "creado": a.creado.isoformat() if a.creado else None,
            "documentos": docs_by_aviso.get(a.id, []),
        }
        for a in rows
    ]


@router.patch("/avisos/{aviso_id}/leer")
def marcar_aviso_leido(aviso_id: int, db: Session = Depends(get_db), token=Depends(_require_portal)):
    """Cliente marca un aviso como leído."""
    rol = token.get("rol", "")
    cliente_ref = (token.get("cliente_ref") or "").strip()

    aviso = db.query(models.AvisoCliente).filter_by(id=aviso_id).first()
    if not aviso:
        raise HTTPException(404, "Aviso no encontrado")
    if rol != "admin" and aviso.cliente_ref != cliente_ref:
        raise HTTPException(403, "Sin acceso a este aviso")

    aviso.leido = True
    db.commit()
    return {"ok": True}


@router.delete("/avisos/{aviso_id}")
def eliminar_aviso(aviso_id: int, db: Session = Depends(get_db), token=Depends(verify_token)):
    """Admin elimina un aviso y sus documentos adjuntos."""
    if token.get("rol") != "admin":
        raise HTTPException(status_code=403, detail="Solo administradores pueden eliminar avisos")
    aviso = db.query(models.AvisoCliente).filter_by(id=aviso_id).first()
    if not aviso:
        raise HTTPException(404, "Aviso no encontrado")
    _borrar_docs_asociados(db, ["aviso"], aviso_id)
    db.delete(aviso)
    db.commit()
    return {"ok": True}


# ─── REPORTES DEL CLIENTE ─────────────────────────────────────────────────────

@router.get("/reportes")
def portal_reporte(
    mes: str = "",
    anio: str = "",
    formato: str = "excel",
    db: Session = Depends(get_db),
    token=Depends(_require_portal),
):
    """Descarga reporte Excel o PDF de afiliados del cliente para un período."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    rol         = token.get("rol", "")
    cliente_ref = (token.get("cliente_ref") or "").strip()

    if rol != "admin" and not cliente_ref:
        raise HTTPException(400, "Usuario sin cliente asociado")

    # ── Afiliados del cliente ─────────────────────────────────────────────────
    query = db.query(models.Afiliado).filter(models.Afiliado.activo == True)
    if rol != "admin":
        query = query.filter(
            models.Afiliado.cliente_txt == cliente_ref,
            models.Afiliado.cliente_txt != None,
            models.Afiliado.cliente_txt != "",
        )
    afiliados = query.order_by(models.Afiliado.nombre).limit(2000).all()

    # ── Facturas del período para cada afiliado ───────────────────────────────
    PAGADOS = ("pagado", "planilla_pagada")

    reporte = []
    for a in afiliados:
        fq = db.query(models.Factura).filter_by(doc=a.doc, afiliado_eliminado=False)
        if mes:
            fq = fq.filter(models.Factura.mes == mes)
        if anio:
            fq = fq.filter(models.Factura.anio == anio)
        facturas = fq.order_by(models.Factura.id.desc()).all()

        try:
            srvs = json.loads(a.servicios or "[]")
        except Exception:
            srvs = []

        if facturas:
            for f in facturas:
                reporte.append({
                    "nombre": a.nombre, "doc": a.doc, "tipo_doc": a.tipo_doc,
                    "empresa": a.empresa, "cargo": a.cargo or "",
                    "eps": a.eps or "", "afp": a.afp or "", "arl": a.arl or "", "ccf": a.ccf or "",
                    "servicios": ", ".join(srvs),
                    "estado_afil": a.estado,
                    "periodo": f"{f.mes} {f.anio}" if f.mes else (anio or ""),
                    "codigo": f.codigo or "",
                    "estado_factura": "Pagada" if f.estado in PAGADOS else f.estado.capitalize(),
                    "valor": f.costos or 0,
                    "banco": f.banco or "",
                    "fecha_pago": f.pagado_en.strftime("%Y-%m-%d") if f.pagado_en else "",
                    "sin_factura": False,
                })
        else:
            reporte.append({
                "nombre": a.nombre, "doc": a.doc, "tipo_doc": a.tipo_doc,
                "empresa": a.empresa, "cargo": a.cargo or "",
                "eps": a.eps or "", "afp": a.afp or "", "arl": a.arl or "", "ccf": a.ccf or "",
                "servicios": ", ".join(srvs),
                "estado_afil": a.estado,
                "periodo": f"{mes} {anio}".strip() if (mes or anio) else "Sin período",
                "codigo": "", "estado_factura": "Sin factura", "valor": 0,
                "banco": "", "fecha_pago": "", "sin_factura": True,
            })

    # ── Totales ───────────────────────────────────────────────────────────────
    total_afil    = len(afiliados)
    con_factura   = [r for r in reporte if not r["sin_factura"]]
    sin_factura   = [r for r in reporte if r["sin_factura"]]
    total_pagado  = sum(r["valor"] for r in con_factura if r["estado_factura"] == "Pagada")
    total_pendiente = sum(r["valor"] for r in con_factura if r["estado_factura"] != "Pagada")
    periodo_label = f"{mes} {anio}".strip() if (mes or anio) else "Todos los períodos"
    fecha_gen     = datetime.now().strftime("%d/%m/%Y %H:%M")
    cliente_label = cliente_ref if rol != "admin" else "Administrador"

    if formato == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte Afiliados"

        # Estilos
        hdr_font  = Font(bold=True, color="FFFFFF", size=11)
        hdr_fill  = PatternFill("solid", fgColor="1B3A6B")
        sub_fill  = PatternFill("solid", fgColor="E8F0FA")
        ok_fill   = PatternFill("solid", fgColor="D1FAE5")
        pend_fill = PatternFill("solid", fgColor="FEE2E2")
        no_fill   = PatternFill("solid", fgColor="FEF9C3")
        thin      = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        )
        center    = Alignment(horizontal="center", vertical="center")
        money_fmt = '#,##0'

        # Título
        ws.merge_cells("A1:N1")
        ws["A1"] = f"Reporte de Afiliados — {cliente_label} — {periodo_label}"
        ws["A1"].font = Font(bold=True, size=14, color="1B3A6B")
        ws["A1"].alignment = center

        ws.merge_cells("A2:N2")
        ws["A2"] = f"Generado: {fecha_gen}  |  Total afiliados: {total_afil}  |  Con factura: {len(con_factura)}  |  Sin factura: {len(sin_factura)}"
        ws["A2"].font = Font(italic=True, size=10, color="555555")
        ws["A2"].alignment = center

        # Resumen financiero
        ws.append([])
        ws.merge_cells("A3:N3")
        ws["A3"] = f"Total pagado: ${total_pagado:,.0f}   |   Total pendiente: ${total_pendiente:,.0f}"
        ws["A3"].font = Font(bold=True, size=11)
        ws["A3"].alignment = center

        ws.append([])  # fila 4 vacía

        # Cabecera
        headers = ["Nombre", "Documento", "Tipo Doc", "Empresa", "Cargo",
                   "EPS", "AFP", "ARL", "CCF", "Servicios",
                   "Estado Afiliado", "Período", "Código", "Estado Factura",
                   "Valor ($)", "Banco", "Fecha Pago"]
        ws.append(headers)
        hdr_row = ws.max_row
        for col, _ in enumerate(headers, 1):
            cell = ws.cell(row=hdr_row, column=col)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = center
            cell.border = thin

        # Filas de datos
        for r in reporte:
            row = [
                r["nombre"], r["doc"], r["tipo_doc"], r["empresa"], r["cargo"],
                r["eps"], r["afp"], r["arl"], r["ccf"], r["servicios"],
                r["estado_afil"], r["periodo"], r["codigo"], r["estado_factura"],
                r["valor"], r["banco"], r["fecha_pago"],
            ]
            ws.append(row)
            dr = ws.max_row
            # Color por estado
            if r["sin_factura"]:
                fill = no_fill
            elif r["estado_factura"] == "Pagada":
                fill = ok_fill
            else:
                fill = pend_fill
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=dr, column=col)
                cell.fill = fill
                cell.border = thin
                if col == 15:  # Valor
                    cell.number_format = money_fmt

        # Anchos de columna
        col_widths = [32, 14, 9, 20, 16, 12, 12, 12, 12, 22, 14, 14, 12, 16, 14, 14, 12]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

        ws.freeze_panes = f"A{hdr_row + 1}"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f"reporte-afiliados-{mes or 'todos'}-{anio or 'todos'}.xlsx"
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="{fname}"'})

    elif formato == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                leftMargin=1.5*cm, rightMargin=1.5*cm,
                                topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        PRIMARY = colors.HexColor("#1B3A6B")
        GREEN   = colors.HexColor("#065F46")
        RED     = colors.HexColor("#991B1B")
        GREEN_BG = colors.HexColor("#D1FAE5")
        RED_BG   = colors.HexColor("#FEE2E2")
        YELLOW_BG= colors.HexColor("#FEF9C3")

        story = []

        # Título
        title_style = ParagraphStyle("title", parent=styles["Heading1"],
                                     textColor=PRIMARY, fontSize=16, spaceAfter=4)
        sub_style   = ParagraphStyle("sub", parent=styles["Normal"],
                                     textColor=colors.grey, fontSize=9, spaceAfter=2)
        story.append(Paragraph(f"Reporte de Afiliados — {cliente_label}", title_style))
        story.append(Paragraph(f"Período: {periodo_label}  |  Generado: {fecha_gen}", sub_style))
        story.append(Spacer(1, 0.3*cm))

        # Resumen
        resumen_data = [
            ["Total afiliados", "Con factura", "Sin factura", "Total pagado", "Total pendiente"],
            [str(total_afil), str(len(con_factura)), str(len(sin_factura)),
             f"${total_pagado:,.0f}", f"${total_pendiente:,.0f}"],
        ]
        res_table = Table(resumen_data, colWidths=[4*cm]*5)
        res_table.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), PRIMARY),
            ("TEXTCOLOR",  (0,0), (-1,0), colors.white),
            ("FONTNAME",   (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE",   (0,0), (-1,-1), 9),
            ("ALIGN",      (0,0), (-1,-1), "CENTER"),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white]),
            ("GRID",       (0,0), (-1,-1), 0.5, colors.grey),
            ("TOPPADDING", (0,0), (-1,-1), 6),
            ("BOTTOMPADDING", (0,0), (-1,-1), 6),
        ]))
        story.append(res_table)
        story.append(Spacer(1, 0.5*cm))

        # Tabla principal — landscape A4 disponible ~26.7cm
        # Columnas: 4.5+2.4+3.2+1.9+1.9+1.9+2.3+2.3+2.6+1.9+1.8 = 26.7cm
        col_hdr = ["Nombre", "Documento", "Empresa", "EPS", "AFP", "ARL",
                   "Período", "Estado Factura", "Valor ($)", "Banco", "Fecha Pago"]
        col_w   = [4.5*cm, 2.4*cm, 3.2*cm, 1.9*cm, 1.9*cm, 1.9*cm,
                   2.3*cm, 2.3*cm, 2.6*cm, 1.9*cm, 1.8*cm]

        cell_style = ParagraphStyle("cell", fontSize=7, leading=9, wordWrap="CJK")

        table_data = [col_hdr]
        row_colors = []
        for i, r in enumerate(reporte, 1):
            table_data.append([
                Paragraph(r["nombre"], cell_style),
                r["doc"],
                Paragraph(r["empresa"] or "", cell_style),
                r["eps"] or "—",
                r["afp"] or "—",
                r["arl"] or "—",
                r["periodo"],
                r["estado_factura"],
                f"${r['valor']:,.0f}",
                r["banco"] or "—",
                r["fecha_pago"] or "—",
            ])
            if r["sin_factura"]:
                row_colors.append(("BACKGROUND", (0,i), (-1,i), YELLOW_BG))
            elif r["estado_factura"] == "Pagada":
                row_colors.append(("BACKGROUND", (0,i), (-1,i), GREEN_BG))
            else:
                row_colors.append(("BACKGROUND", (0,i), (-1,i), RED_BG))

        base_style = [
            ("BACKGROUND",   (0,0), (-1,0), PRIMARY),
            ("TEXTCOLOR",    (0,0), (-1,0), colors.white),
            ("FONTNAME",     (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE",     (0,0), (-1,-1), 7),
            ("ROWBACKGROUNDS",(0,1), (-1,-1), [colors.white]),
            ("GRID",         (0,0), (-1,-1), 0.3, colors.HexColor("#CCCCCC")),
            ("TOPPADDING",   (0,0), (-1,-1), 3),
            ("BOTTOMPADDING",(0,0), (-1,-1), 3),
            ("LEFTPADDING",  (0,0), (-1,-1), 3),
            ("RIGHTPADDING", (0,0), (-1,-1), 3),
            ("ALIGN",        (8,1), (8,-1), "RIGHT"),
            ("VALIGN",       (0,0), (-1,-1), "MIDDLE"),
        ] + row_colors

        main_table = Table(table_data, colWidths=col_w, repeatRows=1)
        story.append(main_table)

        doc.build(story)
        buf.seek(0)
        fname = f"reporte-afiliados-{mes or 'todos'}-{anio or 'todos'}.pdf"
        return StreamingResponse(buf, media_type="application/pdf",
                                 headers={"Content-Disposition": f'attachment; filename="{fname}"'})

    raise HTTPException(400, "Formato inválido. Use 'excel' o 'pdf'")

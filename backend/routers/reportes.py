"""Router de reportes Excel."""
from fastapi import APIRouter, Depends
from sqlalchemy.orm import Session
from database import get_db
import openpyxl
from openpyxl.styles import Font
import crud
import models
from utils.excel_export import hdr_style, xlsx_response
from .deps import verify_token, require_admin

router = APIRouter(prefix="/reportes", tags=["reportes"])


@router.get("/cobro")
def reporte_cobro(
    empresa: str = "", cliente: str = "", tipo: str = "",
    db: Session = Depends(get_db), token=Depends(verify_token)
):
    from datetime import datetime
    from models import COL_TZ
    from const import MESES

    items = crud.get_cobro(db, empresa=empresa, cliente=cliente, tipo=tipo)

    # ── Calcular periodos de mora por afiliado ────────────────────────────────
    docs = list({r["doc"] for r in items if r.get("doc")})
    paid_facts: set = set()
    if docs:
        paid_facts = set(
            (f.doc, f.mes, f.anio)
            for f in db.query(models.Factura.doc, models.Factura.mes, models.Factura.anio)
            .filter(models.Factura.doc.in_(docs))
            .filter(models.Factura.estado.in_(["pagado", "planilla_pagada"]))
            .all()
        )

    cfg = db.query(models.Config).first()
    cfg_corte = (cfg.anio_inicio_cobro, cfg.mes_inicio_cobro) if (
        cfg and cfg.anio_inicio_cobro and cfg.mes_inicio_cobro
    ) else None
    hoy = datetime.now(COL_TZ)

    def _mora(doc: str, fecha_afil: str):
        if not fecha_afil or "-" not in fecha_afil:
            return 0, ""
        try:
            parts = fecha_afil.split("-")
            y0, m0, d0 = int(parts[0]), int(parts[1]), int(parts[2])
        except Exception:
            return 0, ""
        pm, py = m0 + 1, y0
        if pm > 12:
            pm, py = 1, py + 1
        if cfg_corte and (py, pm) < cfg_corte:
            py, pm = cfg_corte[0], cfg_corte[1]
        meses_mora = []
        y, m = py, pm
        while (y, m) <= (hoy.year, hoy.month):
            mes_nom = MESES[m - 1]
            es_actual = (y == hoy.year and m == hoy.month)
            vencido = (not es_actual) or (d0 <= hoy.day)
            if vencido and (doc, mes_nom, str(y)) not in paid_facts:
                meses_mora.append(f"{mes_nom} {y}")
            m += 1
            if m > 12:
                m, y = 1, y + 1
        return len(meses_mora), ", ".join(meses_mora)

    mora_cache: dict = {}
    for r in items:
        doc = r.get("doc")
        if doc and doc not in mora_cache:
            mora_cache[doc] = _mora(doc, r.get("fecha_afiliacion", ""))

    # ── Construir Excel ───────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cobro"
    cols = ["#", "Nombre", "Tipo Doc", "Documento", "Empresa", "Cliente",
            "Día cobro", "Servicios", "Planilla ($)", "Estado",
            "Períodos mora", "Meses en mora"]
    hdr_style(ws, cols)
    for i, r in enumerate(items, 1):
        srvs = r.get("servicios") or []
        srvs_str = ", ".join(srvs) if isinstance(srvs, list) else str(srvs)
        doc = r.get("doc")
        mora_n, mora_str = mora_cache.get(doc, (0, ""))
        ws.append([i, r.get("nombre"), r.get("tipo_doc", ""), doc,
                   r.get("empresa"), r.get("cliente"), r.get("dia_cobro"),
                   srvs_str, r.get("planilla", 0), r.get("estado"),
                   mora_n or "", mora_str])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(col[0].value or "")), 14)
    return xlsx_response(wb, "cobro.xlsx")


@router.get("/afiliados")
def reporte_afiliados(
    q: str = "", estado: str = "", empresa: str = "",
    cliente: str = "", subtipo: str = "",
    db: Session = Depends(get_db), token=Depends(verify_token)
):
    result = crud.get_afiliados(db, q=q, estado=estado, empresa=empresa,
                                cliente=cliente, subtipo=subtipo, skip=0, limit=0)
    items = result.get("items", [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Afiliados"
    cols = ["#", "Nombre", "Tipo Doc", "Documento", "Empresa", "Cliente", "Cargo", "EPS", "ARL", "AFP", "CCF",
            "Subtipo", "Estado", "Estado Servicio", "Servicios", "IBC", "Tel", "Email", "Fecha Ingreso", "Fecha Afiliación"]
    hdr_style(ws, cols)
    for i, a in enumerate(items, 1):
        srvs = a.get("servicios") or []
        srvs_str = ", ".join(srvs) if isinstance(srvs, list) else str(srvs)
        ws.append([i, a.get("nombre"), a.get("tipo_doc", ""), a.get("doc"), a.get("empresa"), a.get("cliente_txt"),
                   a.get("cargo"), a.get("eps"), a.get("arl"), a.get("afp"), a.get("ccf"),
                   a.get("subtipo"), a.get("estado"), a.get("estado_srv"), srvs_str,
                   a.get("ibc"), a.get("tel"), a.get("email"),
                   a.get("fecha_ingreso"), a.get("fecha_afiliacion")])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(col[0].value or "")), 12)
    return xlsx_response(wb, "afiliados.xlsx")


@router.get("/retiros")
def reporte_retiros(
    anio: str = "", mes: str = "",
    db: Session = Depends(get_db), token=Depends(verify_token)
):
    result = crud.get_retiros(db, anio=anio, mes=mes)
    items = result.get("items", []) if isinstance(result, dict) else result
    # Batch lookup tipo_doc desde afiliados (retiro no guarda tipo_doc)
    docs_ret = list({r.get("doc") for r in items if r.get("doc")})
    tipo_doc_map = {}
    cliente_map  = {}
    if docs_ret:
        afils = db.query(models.Afiliado.doc, models.Afiliado.tipo_doc, models.Afiliado.cliente_txt)\
                  .filter(models.Afiliado.doc.in_(docs_ret)).all()
        tipo_doc_map = {a.doc: (a.tipo_doc or "") for a in afils}
        cliente_map  = {a.doc: (a.cliente_txt or "") for a in afils}
        # también buscar en eliminados por si el afiliado ya no existe
        import json as _j
        eliminados_docs = [d for d in docs_ret if d not in tipo_doc_map or d not in cliente_map]
        if eliminados_docs:
            for e in db.query(models.Eliminado).filter(models.Eliminado.doc.in_(eliminados_docs)).all():
                try:
                    datos = _j.loads(e.datos_completos or "{}")
                    if e.doc not in tipo_doc_map:
                        tipo_doc_map[e.doc] = datos.get("tipo_doc", "")
                    if e.doc not in cliente_map:
                        cliente_map[e.doc]  = datos.get("cliente_txt", "")
                except Exception:
                    pass
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Retiros"
    cols = ["#", "Nombre", "Tipo Doc", "Documento", "Empresa", "Cliente", "Fecha", "Motivo", "Mes", "Año", "Observaciones", "Registrado por"]
    hdr_style(ws, cols)
    for i, r in enumerate(items, 1):
        ws.append([i, r.get("nombre"), tipo_doc_map.get(r.get("doc"), ""), r.get("doc"),
                   r.get("empresa"), cliente_map.get(r.get("doc"), ""),
                   r.get("fecha"), r.get("motivo"),
                   r.get("mes"), r.get("anio"), r.get("obs"), r.get("registrado_por")])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(col[0].value or "")), 12)
    return xlsx_response(wb, f"retiros{'_'+anio if anio else ''}{'_'+mes if mes else ''}.xlsx")


@router.get("/financiero")
def reporte_financiero(
    anio: str = "", mes: str = "", cliente: str = "", estado: str = "", banco: str = "",
    db: Session = Depends(get_db), token=Depends(require_admin)
):
    result = crud.get_facturas(db, anio=anio, mes=mes, cliente=cliente,
                               estado=estado, banco=banco, skip=0, limit=0)
    items = result.get("items", []) if isinstance(result, dict) else result

    # Mapa doc → (empresa, subtipo) para enriquecer el reporte
    docs = list({f.get("doc") for f in items if f.get("doc")})
    afil_map = {}
    if docs:
        afils = db.query(models.Afiliado.doc, models.Afiliado.empresa,
                         models.Afiliado.subtipo, models.Afiliado.tipo_doc)\
                  .filter(models.Afiliado.doc.in_(docs)).all()
        afil_map = {a.doc: (a.empresa or "", a.subtipo or "", a.tipo_doc or "") for a in afils}

    # Ingresos adicionales del período para cuadrar con dashboard
    from sqlalchemy import func as _func
    MESES_NUM = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto",
                 "Septiembre","Octubre","Noviembre","Diciembre"]
    _ia_q = db.query(_func.coalesce(_func.sum(models.IngresoAdicional.valor), 0))
    if anio: _ia_q = _ia_q.filter(models.IngresoAdicional.anio == int(anio))
    if mes and mes in MESES_NUM:
        _ia_q = _ia_q.filter(models.IngresoAdicional.mes == MESES_NUM.index(mes) + 1)
    ing_adic_total = float(_ia_q.scalar() or 0)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturación"
    cols = ["#", "Código", "Afiliado", "Tipo Doc", "Documento", "Empresa", "Subtipo", "Servicios",
            "Cliente", "Mes", "Año", "Período (días)", "Ingresos", "Planilla SS",
            "Utilidad", "Banco", "Estado", "Fecha pago"]
    hdr_style(ws, cols)
    tot_ing_pag = tot_plan_pag = tot_util_pag = 0
    tot_ing_pend = tot_plan_pend = tot_util_pend = 0
    docs_con_factura = set()
    ESTADOS_PAGADO = {"pagado", "planilla_pagada"}
    for i, f in enumerate(items, 1):
        emp, sub, tipo_doc = afil_map.get(f.get("doc"), ("", "", ""))
        servicios_txt = ", ".join(s["servicio"] for s in (f.get("servicios_detalle") or []) if s.get("servicio"))
        _fp = f.get("pagado_en") or ""
        if _fp and len(_fp) >= 10:
            _d = _fp[:10].split("-")
            fecha_pago = f"{_d[2]}/{_d[1]}/{_d[0]}" if len(_d) == 3 else _fp[:10]
        else:
            fecha_pago = ""
        ing  = f.get("ingresos", 0) or 0
        cost = f.get("costos", 0) or 0
        util = f.get("utilidad", 0) or 0
        ws.append([i, f.get("codigo"), f.get("nombre_afiliado"), tipo_doc, f.get("doc"),
                   emp, sub, servicios_txt,
                   f.get("cliente"), f.get("mes"), f.get("anio"), f.get("periodo"),
                   ing, cost, util, f.get("banco"), f.get("estado"), fecha_pago])
        if f.get("estado") in ESTADOS_PAGADO:
            tot_ing_pag  += ing;  tot_plan_pag  += cost;  tot_util_pag  += util
        else:
            tot_ing_pend += ing;  tot_plan_pend += cost;  tot_util_pend += util
        if f.get("doc"):
            docs_con_factura.add(f["doc"])

    # Fila total pagadas (facturas)
    last = ws.max_row + 1
    ws.cell(last, 1, "TOTAL PAGADAS (facturas)").font = Font(bold=True, color="166534")
    ws.cell(last, 13, tot_ing_pag).font  = Font(bold=True, color="166534")
    ws.cell(last, 14, tot_plan_pag).font = Font(bold=True, color="166534")
    ws.cell(last, 15, tot_util_pag).font = Font(bold=True, color="166534")
    ws.cell(last, 15).comment = None  # Utilidad bruta por factura (ingresos - costos + conceptos_extra)

    # Fila ingresos adicionales
    last2 = last + 1
    ws.cell(last2, 1, "Ingresos adicionales").font = Font(italic=True, color="1E40AF")
    ws.cell(last2, 13, ing_adic_total).font        = Font(italic=True, color="1E40AF")

    # Fila TOTAL INGRESOS = pagadas + adicionales (= dashboard)
    last3 = last2 + 1
    total_ingresos = tot_ing_pag + ing_adic_total
    fill_total = PatternFill("solid", fgColor="1E40AF")
    for col_idx, val in [(1, "TOTAL INGRESOS (sin descontar nóminas/gastos)"), (13, total_ingresos)]:
        c = ws.cell(last3, col_idx, val)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = fill_total

    # Fila total pendientes
    last4 = last3 + 1
    ws.cell(last4, 1, "TOTAL PENDIENTE").font = Font(bold=True, color="B45309")
    ws.cell(last4, 13, tot_ing_pend).font  = Font(bold=True, color="B45309")
    ws.cell(last4, 14, tot_plan_pend).font = Font(bold=True, color="B45309")
    ws.cell(last4, 15, tot_util_pend).font = Font(bold=True, color="B45309")

    # ── Sección: Afiliados sin factura en el período ──
    afiliados_activos = db.query(models.Afiliado).filter_by(activo=True).order_by(models.Afiliado.nombre).all()
    sin_factura = [a for a in afiliados_activos if a.doc not in docs_con_factura]
    if sin_factura:
        periodo_label = f"{mes} {anio}".strip() if (mes or anio) else "el período"
        sep_row = ws.max_row + 2
        ws.cell(sep_row, 1, f"AFILIADOS SIN FACTURA EN {periodo_label.upper()}").font = Font(bold=True, size=12, color="DC2626")
        hdr_row = sep_row + 1
        sin_cols = ["#", "Nombre", "Tipo Doc", "Documento", "Empresa", "Subtipo", "EPS", "ARL", "Estado"]
        hdr_style(ws, sin_cols, row=hdr_row)
        for j, a in enumerate(sin_factura, 1):
            ws.append([j, a.nombre, a.tipo_doc or "", a.doc, a.empresa or "", a.subtipo or "",
                       a.eps or "", a.arl or "", a.estado_srv or a.estado or ""])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(col[0].value or "")), 14)
    return xlsx_response(wb, f"financiero{'_'+anio if anio else ''}{'_'+mes if mes else ''}.xlsx")


@router.get("/eliminados")
def reporte_eliminados(db: Session = Depends(get_db), token=Depends(verify_token)):
    """Excel con todos los afiliados eliminados."""
    import json as _json

    def _parse(datos):
        try:
            return _json.loads(datos or "{}")
        except Exception:
            return {}

    rows = db.query(models.Eliminado).order_by(models.Eliminado.fecha_eliminacion.desc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Eliminados"
    cols = ["#", "Nombre", "Empresa", "Documento", "EPS", "CCF", "Fecha afiliación",
            "Mes", "Fecha eliminación", "Eliminado por", "Estado planilla"]
    hdr_style(ws, cols)
    ESTADOS = {
        "retiro_pendiente": "Retiro pendiente",
        "planilla_hecha": "Planilla hecha",
        "planilla_pagada": "Planilla pagada",
    }
    for i, r in enumerate(rows, 1):
        d = _parse(r.datos_completos)
        ws.append([
            i,
            r.nombre or "",
            r.empresa or "",
            r.doc or "",
            d.get("eps", "") or "",
            d.get("ccf", "") or "",
            d.get("fecha_afiliacion", "") or "",
            r.mes or "",
            r.fecha_eliminacion or "",
            r.eliminado_por or "",
            ESTADOS.get(r.estado_planilla or "", "—"),
        ])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(col[0].value or "")), 12)
    return xlsx_response(wb, "eliminados.xlsx")

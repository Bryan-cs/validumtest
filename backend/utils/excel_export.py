"""Helpers compartidos para exportaciones Excel."""
import io
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def hdr_style(ws, cols: list, row: int = 1, fg_color: str = "1E40AF"):
    fill = PatternFill("solid", fgColor=fg_color)
    font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(bottom=Side(style="medium", color="FFFFFF"))
    for i, col in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=col)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(horizontal="center")
        c.border = border


def xlsx_response(wb: openpyxl.Workbook, filename: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
